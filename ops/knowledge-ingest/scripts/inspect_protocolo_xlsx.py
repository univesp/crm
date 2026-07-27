#!/usr/bin/env python3
"""Inspect protocol Excel structure (no PII in output)."""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def mask(value: object) -> str:
    text = str(value or "").strip()
    if len(text) <= 4:
        return "***"
    return f"{text[:3]}…{text[-2:]}" if len(text) > 8 else f"{text[:2]}…"


def main() -> int:
    path = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/ProtocoloStatus_Atualizado_ate_202606.xlsx")
    wb = load_workbook(path, read_only=True, data_only=True)
    print("sheets:", wb.sheetnames)
    for sheet_name in wb.sheetnames:
        print("\n=== sheet:", sheet_name, "===")
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        cols = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(header)]
        print("column_count:", len(cols))
        for i, name in enumerate(cols):
            print(f"  [{i}] {name}")

        subject_idx = None
        for i, name in enumerate(cols):
            lower = name.lower()
            if any(token in lower for token in ("assunto", "subject", "titulo", "título", "motivo", "descricao", "descrição", "demanda", "protocolo")):
                subject_idx = i
                break

        total = 0
        subject_counter: Counter[str] = Counter()
        for row in rows:
            total += 1
            if subject_idx is not None and subject_idx < len(row):
                subject = str(row[subject_idx] or "").strip()
                if subject:
                    subject_counter[subject] += 1

        print("data_rows:", total)
        print("subject_column:", cols[subject_idx] if subject_idx is not None else "NOT FOUND")
        print("unique_subjects:", len(subject_counter))
        print("top_10_subjects:")
        for subject, count in subject_counter.most_common(10):
            print(f"  {count:6d}  {mask(subject)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
