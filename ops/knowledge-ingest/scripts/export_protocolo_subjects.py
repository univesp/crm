#!/usr/bin/env python3
"""Export FAQ subject counts from protocol status Excel (no PII output)."""

from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_SHEET = "PROTOCOLOS_STATUS"
DEFAULT_COLUMN = "Motivo_Normalizado"


def export_subjects(
    xlsx_path: Path,
    output_csv: Path,
    sheet_name: str = DEFAULT_SHEET,
    column: str = DEFAULT_COLUMN,
) -> dict[str, int]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Aba '{sheet_name}' não encontrada. Abas: {wb.sheetnames}")
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    cols = [str(c).strip() if c is not None else "" for c in header]
    try:
        col_idx = cols.index(column)
    except ValueError as exc:
        raise ValueError(f"Coluna '{column}' não encontrada. Colunas: {cols}") from exc

    counter: Counter[str] = Counter()
    skipped = 0
    for row in rows:
        if col_idx >= len(row):
            skipped += 1
            continue
        subject = str(row[col_idx] or "").strip()
        if not subject or subject.lower() in {"none", "nan", "(não informado)", "(nao informado)"}:
            skipped += 1
            continue
        counter[subject] += 1

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["assunto", "count"])
        for subject, count in counter.most_common():
            writer.writerow([subject, count])

    return {
        "rows_read": sum(counter.values()) + skipped,
        "subjects_used": sum(counter.values()),
        "unique_subjects": len(counter),
        "skipped": skipped,
        "output": str(output_csv),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Export assuntos de protocolos XLSX para CSV")
    parser.add_argument("xlsx", help="Arquivo ProtocoloStatus*.xlsx")
    parser.add_argument("--output", "-o", help="CSV de saída (assunto,count)")
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--column", default=DEFAULT_COLUMN)
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    output = Path(args.output) if args.output else xlsx_path.with_suffix(".subjects.csv")
    stats = export_subjects(xlsx_path, output, sheet_name=args.sheet, column=args.column)
    print(
        f"OK: {stats['unique_subjects']} assuntos únicos · "
        f"{stats['subjects_used']} protocolos · CSV: {stats['output']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
